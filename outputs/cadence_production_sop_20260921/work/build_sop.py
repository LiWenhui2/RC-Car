from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.dml import MSO_THEME_COLOR_INDEX
from pathlib import Path
import re

OUT = Path(r"H:\RC-Car\outputs\cadence_production_sop_20260921")
OUT.mkdir(parents=True, exist_ok=True)
DOCX = OUT / "RC-CAR-V1.0_Cadence完整设计生产SOP.docx"

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "F4F7FB"
PALE_GOLD = "FFF2CC"
GOLD = "7A5A00"
PALE_RED = "FCE4D6"
RED = "9B1C1C"
PALE_GREEN = "E2F0D9"
GREEN = "2F6B2F"
GRAY = "666666"
LIGHT_GRAY = "F2F4F7"
GRID = "C9D2DC"
WHITE = "FFFFFF"
BLACK = "111111"

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.82)
section.bottom_margin = Inches(0.78)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)
section.header_distance = Inches(0.38)
section.footer_distance = Inches(0.38)

def set_font(run, ascii_name="Calibri", east_asia="Microsoft YaHei", size=None,
             color=None, bold=None, italic=None):
    run.font.name = ascii_name
    if run._element.get_or_add_rPr().rFonts is None:
        run._element.get_or_add_rPr().append(OxmlElement("w:rFonts"))
    rfonts = run._element.get_or_add_rPr().rFonts
    rfonts.set(qn("w:ascii"), ascii_name)
    rfonts.set(qn("w:hAnsi"), ascii_name)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_table_borders(table, color=GRID, size=4):
    tblPr = table._tbl.tblPr
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def set_table_geometry(table, widths_inches, indent_dxa=120):
    total_dxa = sum(round(w * 1440) for w in widths_inches)
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(total_dxa))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), str(indent_dxa))
    tblInd.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    dxa_widths = [round(w * 1440) for w in widths_inches]
    for dxa in dxa_widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(dxa))
        grid.append(gc)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(dxa_widths[min(i, len(dxa_widths)-1)]))
            tcW.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_inches[min(i, len(widths_inches)-1)])
            cell_margins(cell)

def cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    trPr.append(node)

def keep_with_next(p):
    p.paragraph_format.keep_with_next = True

def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, size=9, color=GRAY)
    fldChar1 = OxmlElement("w:fldChar"); fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText"); instrText.set(qn("xml:space"), "preserve"); instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)
    r2 = paragraph.add_run(" 页")
    set_font(r2, size=9, color=GRAY)

def setup_styles():
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.2)
    pf = normal.paragraph_format
    pf.space_after = Pt(4)
    pf.line_spacing = 1.15
    pf.widow_control = True
    for name, size, color, before, after in [
        ("Heading 1",16,BLUE,14,8),
        ("Heading 2",13,BLUE,10,5),
        ("Heading 3",11.5,DARK_BLUE,8,4),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    for name in ["List Bullet", "List Number"]:
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(10.2)
        st.paragraph_format.left_indent = Inches(0.375)
        st.paragraph_format.first_line_indent = Inches(-0.188)
        st.paragraph_format.space_after = Pt(2)
        st.paragraph_format.line_spacing = 1.15
    if "Procedure Step" not in styles:
        st = styles.add_style("Procedure Step", WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = styles["Normal"]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(10.2)
        st.paragraph_format.left_indent = Inches(0.375)
        st.paragraph_format.first_line_indent = Inches(-0.188)
        st.paragraph_format.space_after = Pt(2)
        st.paragraph_format.line_spacing = 1.15
    if "Code Block" not in styles:
        st = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        st.font.name = "Consolas"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(8.5)
        st.paragraph_format.left_indent = Inches(0.18)
        st.paragraph_format.right_indent = Inches(0.18)
        st.paragraph_format.space_before = Pt(2)
        st.paragraph_format.space_after = Pt(2)
        st.paragraph_format.line_spacing = 1.0

setup_styles()

# Header/footer
header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hr = hp.add_run("RC-CAR-V1.0  |  Cadence完整设计与生产SOP")
set_font(hr, size=8.5, color=GRAY, bold=True)
footer = section.footer
fp = footer.paragraphs[0]
page_number(fp)

def add_heading(text, level=1):
    p = doc.add_heading(text, level=level)
    keep_with_next(p)
    return p

def add_para(text="", bold_prefix=None, italic=False, color=None, align=None,
             after=6, keep=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_font(r1, size=10.2, color=color or BLACK, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_font(r2, size=10.2, color=color or BLACK, italic=italic)
    else:
        r = p.add_run(text)
        set_font(r, size=10.2, color=color or BLACK, italic=italic)
    if keep:
        keep_with_next(p)
    return p

def add_bullets(items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375 + level*0.22)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        r = p.add_run(item)
        set_font(r, size=10.2)

def add_steps(items):
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph(style="Procedure Step")
        r1 = p.add_run(f"{i}. ")
        set_font(r1, size=10.2, color=DARK_BLUE, bold=True)
        r2 = p.add_run(item)
        set_font(r2, size=10.2)

def add_checklist(items, title=None):
    if title:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(title)
        set_font(r, size=10.2, color=DARK_BLUE, bold=True)
        keep_with_next(p)
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run("☐  " + item)
        set_font(r, size=10.2)

def add_code(lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [6.5])
    set_table_borders(table, color="D5DCE5", size=4)
    cell = table.cell(0,0)
    shade(cell, LIGHT_GRAY)
    p = cell.paragraphs[0]
    p.style = doc.styles["Code Block"]
    r = p.add_run(lines)
    set_font(r, ascii_name="Consolas", east_asia="Microsoft YaHei", size=8.5, color="243447")

def add_callout(label, text, kind="note"):
    fill, color = {
        "note": (PALE_BLUE, DARK_BLUE),
        "warn": (PALE_GOLD, GOLD),
        "risk": (PALE_RED, RED),
        "pass": (PALE_GREEN, GREEN),
    }[kind]
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [6.5])
    set_table_borders(table, color=color, size=6)
    c = table.cell(0,0)
    shade(c, fill)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run(label + "：")
    set_font(r1, size=9.8, color=color, bold=True)
    r2 = p.add_run(text)
    set_font(r2, size=9.8, color=BLACK)

def add_table(headers, rows, widths, font_size=9.2, header_fill=NAVY):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_geometry(table, widths)
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    cant_split(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        shade(c, header_fill)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(h))
        set_font(r, size=font_size, color=WHITE, bold=True)
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        cant_split(table.rows[-1])
        for i, value in enumerate(row):
            c = cells[i]
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2 == 1:
                shade(c, "F8FAFC")
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            r = p.add_run(str(value))
            set_font(r, size=font_size)
    return table

def add_stage_intro(number, title, goal, inputs, outputs, gate):
    add_heading(f"阶段 {number}：{title}", 1)
    rows = [
        ("目标", goal),
        ("输入", inputs),
        ("输出", outputs),
        ("放行条件", gate),
    ]
    table = add_table(["项目", "内容"], rows, [1.15, 5.35], font_size=9.5, header_fill=BLUE)
    for row in table.rows[1:]:
        shade(row.cells[0], LIGHT_BLUE)
        for run in row.cells[0].paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(DARK_BLUE)

def page_break():
    doc.add_page_break()

# Cover page: editorial_cover pattern
for _ in range(4):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("工程操作手册 · SOP")
set_font(r, size=11, color=GOLD, bold=True)
p.paragraph_format.space_after = Pt(18)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("RC-CAR-V1.0")
set_font(r, size=30, color=NAVY, bold=True)
p.paragraph_format.space_after = Pt(8)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("从元器件建库、原理图、PCB到打样装车的完整流程")
set_font(r, size=16, color=DARK_BLUE, bold=True)
p.paragraph_format.space_after = Pt(8)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("适用工具：OrCAD X Capture · Padstack Editor · OrCAD/Allegro PCB Editor")
set_font(r, size=10.5, color=GRAY)
p.paragraph_format.space_after = Pt(36)
add_callout("项目基线", "P1采用模块载板：ESP32-C3 SuperMini与TB6612FNG模块插接；车辆板优先，遥控器板随后；第二驱动和前轮电机接口保留但首版可不装。", "note")
add_callout("安全边界", "MG995和电机主电流不得经过ESP32或细杜邦线；USB 5V不得与外部5V硬并联；任何未冻结的连接器、保护器件和安装孔都不得直接投板。", "risk")
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("版本：1.0  |  日期：2026-09-21  |  状态：项目执行基线")
set_font(r, size=9.5, color=GRAY)

page_break()
add_heading("如何使用本手册", 1)
add_para("本手册不是一次性阅读材料，而是后续逐阶段实施的项目主线。每完成一个阶段，就保存对应的Cadence工程文件、检查记录和版本快照；只有满足该阶段“放行条件”后，才进入下一阶段。")
add_callout("菜单差异说明", "OrCAD X 24.1/25.1、传统OrCAD PCB Editor以及Allegro PCB Editor的菜单名称可能略有不同。本文同时给出核心对象和常见入口；如果界面文字不同，应优先按“对象/目的”查找，不要仅机械匹配菜单。", "warn")
add_table(["角色", "主要任务", "本项目负责人"], [
    ("电气设计", "冻结电路、引脚、网络和保护方案；完成ERC", "用户/设计者"),
    ("库管理员", "制作并验证符号、Padstack和PCB封装", "用户/设计者"),
    ("PCB设计", "约束、布局、布线、DRC、生产输出", "用户/设计者"),
    ("装配测试", "来料检查、焊接、上电、固件和整车验收", "用户/设计者"),
], [1.15, 4.3, 1.05])

add_heading("全过程路线图", 2)
add_table(["阶段", "名称", "关键产物", "禁止越过的门槛"], [
    ("0-2", "准备与规格冻结", "冻结表、实测记录、最终BOM", "模块针序/连接器/孔位未确认"),
    ("3-5", "Cadence建库", ".pad / .dra / .psm / .olb", "未做1:1实物套孔"),
    ("6-8", "原理图与设计同步", ".dsn / ERC / netlist", "Pin-Pad编号不一致或ERC未闭环"),
    ("9-12", "PCB约束、布局、布线", ".brd / DRC / 3D检查", "电流规则未核算或DRC未清零"),
    ("13-14", "生产文件与下单", "Gerber/钻孔/BOM/坐标", "独立Gerber复核未通过"),
    ("15-17", "焊接、上电、整车验收", "测试记录/问题闭环", "上一上电阶段异常未排除"),
], [0.75, 1.55, 2.1, 2.1])

add_heading("目录", 2)
toc = [
    "1. 项目基线与固定约束", "2. 阶段0：资料备份与工程准备", "3. 阶段1：需求、电源与机械冻结",
    "4. 阶段2：BOM与器件选型冻结", "5. 阶段3：Cadence工程和库路径", "6. 阶段4：Padstack制作",
    "7. 阶段5：PCB封装与原理图符号", "8. 阶段6：车辆板原理图", "9. 阶段7：遥控器原理图",
    "10. 阶段8：ERC、封装绑定与设计同步", "11. 阶段9：板框、叠层和约束", "12. 阶段10：布局",
    "13. 阶段11：布线与覆铜", "14. 阶段12：DRC、评审与冻结", "15. 阶段13：生产文件输出",
    "16. 阶段14：Gerber复核与PCB下单", "17. 阶段15：来料检查与焊接", "18. 阶段16：分阶段上电与固件",
    "19. 阶段17：整车验收与P2闭环", "附录A-D：网络、文件、检查表与资料来源"
]
add_bullets(toc)

add_heading("1. 项目基线与固定约束", 1)
add_heading("1.1 V1硬件架构", 2)
add_bullets([
    "车辆结构：四轮底盘，前轮由MG995舵机转向，后轮由两只TT减速电机驱动。V1不采用四轮差速转向。",
    "车辆主控：ESP32-C3 SuperMini开发板，P1通过2×8排母插接，不在首板改为裸ESP32芯片。",
    "电机驱动：TB6612FNG红色模块U2控制两个后轮；U3和两个前轮电机接口仅作未来扩展预留。",
    "遥控器：第二块ESP32-C3 SuperMini读取PS2双轴摇杆，以ESP-NOW发送throttle、steering、buttons和sequence。",
    "安全策略：通信超过100 ms无有效包时，PWM清零、STBY拉低、舵机回安全中位，并要求摇杆回中后重新解锁。",
])
add_heading("1.2 已验证GPIO基线", 2)
add_table(["功能网络", "车辆端GPIO", "连接对象", "状态"], [
    ("LEFT_IN1", "GPIO3", "U2/U3 AIN1", "已验证"),
    ("LEFT_IN2", "GPIO4", "U2/U3 AIN2", "已验证"),
    ("LEFT_PWM", "GPIO5", "U2/U3 PWMA", "已验证"),
    ("MOTOR_STBY", "GPIO6", "U2/U3 STBY", "已验证"),
    ("RIGHT_IN1", "GPIO7", "U2/U3 BIN1", "已验证"),
    ("RIGHT_IN2", "GPIO10", "U2/U3 BIN2", "已验证"),
    ("RIGHT_PWM", "GPIO20", "U2/U3 PWMB", "已验证"),
    ("SERVO_PWM", "GPIO21", "J3 MG995信号", "已验证"),
    ("BAT_ADC", "GPIO0", "R1/R2/C7", "DNP预留"),
    ("I2C_SDA / I2C_SCL", "GPIO1 / GPIO2", "J8扩展", "DNP预留；GPIO2为启动相关脚"),
], [1.55, 1.15, 2.3, 1.5])
add_callout("必须保持", "GPIO18和GPIO19保留给原生USB；GPIO8和GPIO9为启动相关引脚，不用于当前车辆控制。通信层始终保持正油门代表前进，电机安装方向由执行层的MOTOR_A_DIRECTION=-1、MOTOR_B_DIRECTION=-1修正。", "note")

add_heading("1.3 电源架构", 2)
add_table(["网络", "来源", "负载", "设计要求"], [
    ("VBAT_RAW", "4×AA镍氢电池", "保险丝/反接保护入口", "标称4.8 V，充满约5.2-5.6 V"),
    ("VBAT_PROTECTED", "F1/Q1之后", "TB6612 VM、外置5V模块输入", "大电流、短路径、宽铜皮"),
    ("SERVO_VBAT", "受保护电池大电流分支", "MG995红线", "不得经过ESP32或细排针"),
    ("REG_5V", "板外固定5V升降压模块", "经D1送ESP32 5V", "与USB 5V隔离，严禁硬并联"),
    ("3V3_LOGIC", "SuperMini板载稳压输出", "TB6612 VCC、逻辑接口", "不得给舵机供电"),
], [1.25, 1.65, 1.75, 1.85])
add_callout("当前V1决定", "MG995直接由4节AA镍氢电池的受保护大电流分支供电；现有标称1 A的5V模块只负责ESP32逻辑供电，不负责舵机。", "warn")

add_stage_intro("0", "资料备份与工程准备", "建立可追溯、可回退的设计环境。", "现有接线说明、软件、采购表、实物模块、Cadence安装", "工程目录、版本规则、初始备份", "源资料已归档；任何后续输出都能追溯到版本")
add_heading("2.1 建立工程目录", 2)
add_code("RC-Car/\n  cadence/\n    library/symbols/\n    library/padstacks/\n    library/footprints/\n    library/step/\n    vehicle/schematic/\n    vehicle/pcb/\n    controller/schematic/\n    controller/pcb/\n  manufacturing/vehicle/P1/\n    gerber/ drill/ assembly/ bom/ reports/\n  docs/\n  test-records/")
add_para("不要把自建库放进Cadence安装目录。安装升级或重装可能覆盖系统目录，而且项目移交时无法保证其他电脑存在相同路径。")
add_heading("2.2 版本与备份规则", 2)
add_bullets([
    "正式工程文件名使用ASCII字符，避免老版本工具或制造插件对中文路径兼容不一致。示例：RC_CAR_V1_VEHICLE_P1.opj、RC_CAR_V1_VEHICLE_P1.brd。",
    "关键里程碑另存快照：P1_01_LIB、P1_02_SCH_ERC、P1_03_PLACEMENT、P1_04_ROUTED、P1_05_RELEASE。",
    "Gerber发布后不得覆盖原发布目录；任何修改都生成P2或P1R1，并保留变更说明。",
    "每次进入下一阶段前，将工程、库和检查报告整体打包备份。",
])
add_checklist([
    "Cadence版本、许可证类型和安装日期已记录。",
    "项目目录不位于临时目录、下载目录或同步冲突目录。",
    "现有固件、接线说明和采购表已备份。",
    "建立test-records用于保存万用表、电流和上电结果。",
], "阶段0检查")

page_break()
add_stage_intro("1", "需求、电源与机械冻结", "把影响电路和封装的未知项转换成可测量、可签字的输入。", "实物模块、底盘、卡尺、万用表、数据手册", "机械尺寸表、电流测试计划、电源框图、P0问题关闭记录", "5个P0项目全部确认；板框和接口方向可确定")
add_heading("3.1 必须关闭的P0项目", 2)
add_table(["对象", "必须确认", "方法", "放行标准"], [
    ("ESP32-C3 SuperMini", "顶视针序、两排间距、排针直径、USB/天线方向", "卡尺+丝印+万用表", "符号、封装、实物一致"),
    ("TB6612红色模块", "顶视针序、排间距、VM/VCC方向", "卡尺+丝印+通断复核", "U2/U3无需掰针即可插拔"),
    ("2P电池/电机端子", "具体料号、孔径、针距、开口、电流", "数据手册+实物", "封装由具体料号驱动"),
    ("舵机3P接口", "直针或防呆座、GND/PWR/SIG顺序", "MG995插头实配", "插头不能反插或丝印无歧义"),
    ("底盘安装孔", "孔中心距、铜柱外径、周边净空", "卡尺测量", "4孔与70×50 mm初始板框不干涉"),
], [1.35, 2.2, 1.45, 1.5])
add_heading("3.2 实测记录模板", 2)
add_table(["项目", "标称/资料值", "实测值", "量具/条件", "复核人/日期"], [
    ("SuperMini外形", "约22.5×18.0 mm", "________", "卡尺", "________"),
    ("SuperMini排间距", "初值15.24 mm", "________", "卡尺/套孔纸", "________"),
    ("TB6612外形", "约18.3×21.1 mm", "________", "卡尺", "________"),
    ("TB6612排间距", "初值15.24 mm", "________", "卡尺/套孔纸", "________"),
    ("底盘孔中心距", "建议初值62×42 mm", "________", "卡尺", "________"),
    ("后轮电机启动/堵转电流", "未知", "________", "限流电源/电流表", "________"),
    ("MG995峰值电流", "未知", "________", "实际负载", "________"),
], [1.6, 1.45, 1.15, 1.25, 1.05])
add_callout("不要猜", "网络商品图只能用于形成初值，不能替代实物复核。模块克隆板的排间距、针序、USB位置和板边尺寸可能变化。", "risk")

page_break()
add_stage_intro("2", "BOM与器件选型冻结", "把“功能描述”变成具体可采购、可建库的器件。", "Cadence建库BOM、实物尺寸、数据手册、电流测试结果", "冻结BOM、DNP清单、替代料规则、最终封装清单", "每个实装位号都有具体料号或确定模块；每个DNP位号有明确原因")
add_heading("4.1 车辆板P1位号", 2)
add_table(["位号", "器件", "P1状态", "关键说明"], [
    ("U1", "ESP32-C3 SuperMini", "安装", "2×8排母；天线朝板边"),
    ("U2", "TB6612FNG模块", "安装", "后左/后右电机"),
    ("U3", "TB6612FNG模块", "DNP预留", "未来前轮电机；控制网络与U2共享"),
    ("J1", "电池输入2P", "安装", "额定电流由实测决定"),
    ("J2", "外置5V模块4P", "安装", "IN+/GND/5V/GND"),
    ("J3", "MG995舵机3P", "安装", "1 GND / 2 SERVO_VBAT / 3 PWM"),
    ("J4/J5", "后左/后右电机2P", "安装", "开口朝板边"),
    ("J6/J7", "前左/前右电机2P", "DNP预留", "U3未装时无输出"),
    ("J8", "I2C扩展4P", "可DNP", "3V3/GND/SDA/SCL"),
    ("C1/C2", "1000 uF / ≥10 V电解", "安装", "VM和舵机入口储能"),
    ("R1/R2/C7", "100k/100k/100nF", "DNP预留", "BAT_ADC分压与滤波"),
    ("F1", "保险丝或PTC", "建议安装", "需按电机/舵机电流选型"),
    ("Q1/R3", "PMOS反接保护", "方案冻结后安装", "料号、封装、引脚和温升需验证"),
    ("D1", "肖特基二极管", "建议安装", "REG_5V到ESP_5V防回灌"),
    ("TP1-TP6", "测试点", "安装", "VBAT/5V/3V3/GND/STBY/SERVO"),
    ("H1-H4", "M3安装孔", "安装", "NPTH 3.2 mm，孔位实测冻结"),
], [0.85, 2.2, 1.15, 2.3], font_size=8.9)
add_heading("4.2 遥控器板位号", 2)
add_table(["位号", "器件", "状态", "关键说明"], [
    ("U101", "ESP32-C3 SuperMini", "安装", "GPIO0/1/2连接摇杆"),
    ("J101", "PS2摇杆5P接口", "安装", "3.3V供电，针序按实物"),
    ("J102", "0.96英寸I2C OLED 4P", "可选", "不同模块VCC/GND顺序可能不同"),
    ("SW101/SW102", "功能按键", "可选", "摇杆SW已承担急停"),
    ("D101/R101", "状态LED/限流电阻", "可选", "按亮度调整阻值"),
    ("R102/R103/C101/C102", "1k串联+100nF对地", "建议安装", "两路ADC滤波"),
    ("J103", "外部电池/稳压接口", "DNP", "遥控器电源架构冻结后再做"),
], [1.05, 2.45, 1.05, 1.95])
add_checklist([
    "每个器件都有Manufacturer、MPN、Value、Tolerance、Voltage/Current Rating。",
    "所有连接器均有具体料号、配套胶壳、端子和线序。",
    "DNP器件在BOM和原理图属性中统一标记。",
    "极性器件的1脚、正负极或A/K已记录。",
    "没有以“随便一个2P端子”作为最终选型。",
], "阶段2检查")

add_stage_intro("3", "Cadence工程和库路径", "建立稳定、可移交的Cadence工作环境。", "冻结目录、OrCAD X Capture、Padstack Editor、PCB Editor", "工程文件、库路径配置、空白库", "Capture能找到.olb；PCB Editor能找到.pad和.psm")
add_heading("5.1 建立库文件", 2)
add_steps([
    "在cadence/library/symbols中创建项目原理图库，例如RC_CAR_V1.olb。Capture中通常使用File > New > Library，再保存到项目目录。",
    "在cadence/library/padstacks保存所有自建.pad；在footprints保存.dra与.psm；在step保存STEP模型。",
    "启动PCB Editor，打开Setup > User Preferences（不同版本也可能在Preferences）。找到Paths/Library，将padpath加入padstacks目录，将psmpath加入footprints目录。",
    "把项目库路径放在系统库路径前或明确记录搜索顺序，避免同名封装被系统库覆盖。",
    "新建一个测试板，尝试调用一个自建Padstack和一个封装，确认路径有效。",
])
add_heading("5.2 命名规范", 2)
add_table(["对象", "命名示例", "规则"], [
    ("Padstack", "PTH_R1P90_D1P00", "单位和尺寸写入名称；PTH/NPTH/SMD明确"),
    ("封装", "MOD_ESP32C3_SM_2X08_P2P54_R15P24", "器件类型+针数+间距+关键机械尺寸"),
    ("原理图符号", "ESP32_C3_SUPERMINI_16P", "体现模块与引脚数量"),
    ("网络", "VBAT_PROTECTED", "大写、含义明确，不使用NET001"),
    ("板文件", "RC_CAR_VEHICLE_P1.brd", "板卡+版本"),
], [1.1, 2.65, 2.75])
add_callout("常见故障", "封装文件存在但放置列表里找不到，首先检查psmpath；封装能找到但提示Padstack丢失，检查padpath。路径不要指向安装目录。", "note")

add_stage_intro("4", "Padstack制作", "建立可复用、可制造的焊盘定义。", "连接器/模块排针/安装孔数据手册和实测值", ".pad文件和Padstack检查记录", "钻孔、铜盘、阻焊、钢网、镀铜属性均正确")
add_heading("6.1 制作模块排针通孔Padstack", 2)
add_steps([
    "打开Padstack Editor，选择File > New，类型选择Plated Through Hole。单位统一为毫米。",
    "Drill设置为圆孔。初始可用1.00 mm；如果实测排针较粗则用1.10 mm。不要在未测量时放大到过松。",
    "Regular Pad设置外径1.90 mm左右的圆形铜盘；1脚可另建方形或矩形版本，以便视觉识别。",
    "Anti Pad与Thermal Relief按两层板和覆铜需要设置。简单两层板也应确认平面连接方式不是意外断开。",
    "Top/Bottom Solder Mask设置开窗。开窗扩展遵循板厂能力；不要把阻焊开窗做得比铜盘小。",
    "通孔元件一般不需要Paste Mask。确认未意外生成钢网开口。",
    "保存为PTH_R1P90_D1P00.pad和对应方形1脚版本。",
])
add_heading("6.2 其他Padstack", 2)
add_table(["用途", "类型", "关键输入", "注意"], [
    ("5.08 mm螺丝端子", "PTH", "端子脚宽/厚、推荐孔径、铜盘", "必须按具体料号；方脚孔需留装配余量"),
    ("电解电容", "PTH", "脚径、脚距、推荐孔径", "正极标识属于封装，不属于Padstack"),
    ("M3安装孔", "NPTH", "成品孔3.2 mm", "明确Non-plated；周围设置铜避让"),
    ("测试点", "SMD或PTH", "探针直径、阻焊开窗", "不得被模块或连接器遮挡"),
    ("0805电阻电容", "SMD", "IPC/厂商推荐焊盘", "Top Paste与Top Mask均需定义"),
], [1.5, 0.8, 2.15, 2.05])
add_checklist([
    "钻孔为成品孔还是刀具孔的含义已确认。",
    "PTH与NPTH属性没有选反。",
    "通孔Padstack没有不必要的Paste开口。",
    "阻焊开窗和铜盘尺寸满足板厂能力。",
    "文件名不含空格和含糊缩写。",
], "Padstack放行检查")

add_stage_intro("5", "PCB封装与原理图符号", "完成模块、连接器、保护器件和机械对象建库，并验证Pin-Pad一致性。", "Padstack、冻结BOM、数据手册、实物", ".dra/.psm/.olb、1:1套孔记录、Pin-Pad对照表", "所有实装封装完成1:1检查；符号引脚号与焊盘号一致")
add_heading("7.1 PCB封装通用步骤", 2)
add_steps([
    "在PCB Editor中File > New，选择Package Symbol（或Package Symbol Wizard）。单位统一为毫米，原点放在1脚或器件机械中心，并记录选择。",
    "使用Layout > Pins或等效放置命令，调用已验证Padstack。按数据手册坐标逐个放置焊盘，并设置Pin Number。",
    "在Package Geometry / Assembly_Top绘制真实器件本体轮廓；在Silkscreen_Top绘制不压焊盘的装配识别轮廓。",
    "创建Place_Bound_Top或对应放置边界；为USB、插头、螺丝刀和拔插动作预留操作空间。",
    "添加1脚标记、极性、USB方向、VM/VCC方向等关键丝印。丝印不能覆盖焊盘和阻焊开窗。",
    "保存.dra，并执行Create Symbol/生成.psm。查看消息窗口，解决缺失Padstack或边界错误。",
])
add_heading("7.2 SuperMini封装", 2)
add_bullets([
    "两排各8孔，排内2.54 mm；排间初值15.24 mm，最终以实物套孔为准。",
    "本体初值约22.5×18.0 mm；丝印明确USB端、天线端和1脚。",
    "天线端尽量伸出载板边缘。至少在天线正下方及外侧建立禁止铜、禁止走线和禁止放置区域；最终范围以实物天线位置及Espressif建议复核。",
    "预留USB插头插拔空间，不要让螺丝端子或高电解电容挡住USB。",
])
add_heading("7.3 TB6612模块封装", 2)
add_bullets([
    "两排各8孔，排内2.54 mm；排间初值15.24 mm；外形初值18.3×21.1 mm。",
    "丝印必须标出VM、VCC和1脚，避免模块反插或把VM/VCC误认为同类电源。",
    "U2与U3复用同一封装；布局时两者方向尽量一致，便于检查和维修。",
])
add_heading("7.4 建议Pin-Pad编号", 2)
add_table(["模块", "Pad编号", "实物名称", "说明"], [
    ("SuperMini", "1-8", "5V,GND,3V3,GPIO4,GPIO3,GPIO2,GPIO1,GPIO0", "顶视一侧连续编号"),
    ("SuperMini", "9-16", "GPIO20,GPIO21,GPIO5,GPIO6,GPIO7,GPIO8,GPIO9,GPIO10", "从对侧下方向上连续编号"),
    ("TB6612", "1-8", "PWMA,AIN2,AIN1,STBY,BIN1,BIN2,PWMB,GND", "顶视左排"),
    ("TB6612", "9-16", "GND,BO1,BO2,AO2,AO1,GND,VCC,VM", "顶视右排反向连续"),
], [1.05, 0.85, 3.55, 1.05], font_size=8.7)
add_callout("编号原则", "上述编号不是行业唯一方案。可以采用其他编号，但Capture符号、PCB封装和实物方向必须完全一致，并保存对照表。", "warn")

add_heading("7.5 Capture原理图符号", 2)
add_steps([
    "打开RC_CAR_V1.olb，在Part Editor中新建器件。U1/U2/U3使用与封装完全一致的Pin Number。",
    "Pin Name写GPIO、VM、VCC、AO1等真实名称；电源脚、输入、输出和双向脚类型按功能设置，避免ERC误报。",
    "模块符号按功能分组：电源脚置于上/下，控制输入置左，功率输出置右。视觉分组不改变Pin Number。",
    "在属性中准备PCB Footprint、Value、Manufacturer、MPN、DNP、Description等字段。",
    "连接器符号必须显示针脚编号；电池、舵机和电解电容使用清楚的极性标记。",
])
add_heading("7.6 1:1封装检查", 2)
add_steps([
    "从PCB Editor按1:1比例打印封装。打印设置不得选择Fit to Page。",
    "用直尺先检查纸上10 mm标尺是否仍为10 mm，排除打印缩放。",
    "将SuperMini和TB6612实际排针放在纸上套孔；连接器、电容也逐一套孔。",
    "确认模块无需掰针；检查USB、插头、螺丝刀、拔插方向和相邻器件高度。",
    "记录检查人、日期、打印比例和结果；有任何不确定即返回修改封装。",
])

page_break()
add_stage_intro("6", "车辆板原理图", "把已验证接线迁移成清晰、可检查、可制造的车辆载板原理图。", "已验证GPIO、电源架构、符号库、封装库、冻结BOM", ".opj/.dsn、网络表、初版BOM、ERC记录", "所有连接器针序明确；电源路径正确；ERC严重问题为0")
add_heading("8.1 新建工程与页面", 2)
add_steps([
    "启动OrCAD X Capture，File > New > Project，选择PCB类型工程，保存到cadence/vehicle/schematic。",
    "创建4个页面：01_POWER、02_MCU、03_MOTOR_SERVO、04_CONNECTORS_TEST。页面名称保持ASCII，标题栏可写中文说明。",
    "加入RC_CAR_V1.olb。先放置所有器件并确认位号，再开始连线。",
])
add_heading("8.2 页面01_POWER", 2)
add_bullets([
    "J1输入形成VBAT_RAW；经F1和最终确认的反接保护形成VBAT_PROTECTED。",
    "VBAT_PROTECTED分别送U2/U3的VM、J2外置5V模块输入和SERVO_VBAT分支。",
    "J2定义：1 VBAT_PROTECTED/IN+，2 GND/IN-，3 REG_5V/OUT+，4 GND/OUT-。",
    "REG_5V经D1或最终确认的电源隔离方案送U1的5V。USB 5V与外部5V不得硬并联。",
    "C1放在电机VM入口，C2放在舵机接口附近；两者正负极必须明确。",
    "所有地最终共地，但布局布线阶段要避免电机/舵机大电流回流穿过U1地脚附近。",
])
add_heading("8.3 页面02_MCU", 2)
add_table(["U1引脚", "网络", "去向/说明"], [
    ("5V", "ESP_5V", "来自REG_5V隔离后输出"), ("3V3", "3V3_LOGIC", "U2/U3 VCC和I2C"),
    ("GPIO3/4/5", "LEFT_IN1/LEFT_IN2/LEFT_PWM", "U2/U3 A通道共享"),
    ("GPIO6", "MOTOR_STBY", "U2/U3共享；失联时拉低"),
    ("GPIO7/10/20", "RIGHT_IN1/RIGHT_IN2/RIGHT_PWM", "U2/U3 B通道共享"),
    ("GPIO21", "SERVO_PWM", "J3信号"),
    ("GPIO0", "BAT_ADC", "R1/R2/C7 DNP预留"),
    ("GPIO1/2", "I2C_SDA/I2C_SCL", "J8；GPIO2启动相关"),
], [1.45, 2.3, 2.75])
add_heading("8.4 页面03_MOTOR_SERVO", 2)
add_bullets([
    "U2 AO1/AO2接J4后左电机，BO1/BO2接J5后右电机。",
    "U3 AO1/AO2接J6前左电机，BO1/BO2接J7前右电机；U3标DNP。",
    "U2和U3共享LEFT/RIGHT控制网络与MOTOR_STBY。",
    "J3舵机针序固定为1 GND、2 SERVO_VBAT、3 SERVO_PWM；丝印同时标棕/黑、红、橙/黄。",
    "C3-C6可在PCB保留，但100nF抑制电容优先直接焊在电机端子。",
])
add_heading("8.5 页面04_CONNECTORS_TEST", 2)
add_bullets([
    "J8：3V3、GND、SDA、SCL；TP1-TP6：VBAT、5V、3V3、GND、STBY、SERVO。",
    "每个连接器旁放置明确网络标号，不依赖跨页长导线表达。",
    "测试点必须是有封装的实物器件/焊盘，不能只在原理图写文字。",
])
add_checklist([
    "U1/U2/U3的PCB Footprint属性已填写且名称与.psm一致。",
    "所有连接器针脚编号已与线束定义复核。",
    "VBAT_RAW、VBAT_PROTECTED、SERVO_VBAT、REG_5V、ESP_5V、3V3_LOGIC没有混用。",
    "DNP属性在U3、J6/J7可选装件、BAT_ADC预留中一致。",
    "电解电容、D1、Q1、LED等极性清楚。",
], "车辆原理图自检")

page_break()
add_stage_intro("7", "遥控器原理图", "在不改变无线协议的前提下完成摇杆、显示和供电接口。", "U101、摇杆实物针序、可选OLED、遥控器供电决定", "遥控器.opj/.dsn、ERC、BOM", "摇杆ADC范围和供电安全；USB与外部电源不硬并联")
add_heading("9.1 基础连接", 2)
add_table(["对象", "连接", "要求"], [
    ("JOY VRx", "GPIO0", "3.3V供电时输出不超过ADC范围；1k串联+100nF对地"),
    ("JOY VRy", "GPIO1", "同上"),
    ("JOY SW", "GPIO2", "低电平有效；复位/上电时不要按住"),
    ("OLED", "I2C SDA/SCL", "针序按实物；确认3.3V兼容"),
    ("功能按键", "空闲GPIO到GND", "使用上拉；GPIO分配需与固件同步"),
    ("状态LED", "GPIO经限流电阻", "极性和阻值明确"),
], [1.25, 1.85, 3.4])
add_heading("9.2 遥控器供电决策", 2)
add_bullets([
    "原型和P1可直接使用USB供电，此时J103及电池稳压支路DNP。",
    "若需要脱线运行，应确定电池类型、稳压拓扑、开关和充电方式，再独立评审。不得把4节AAA原始电压直接接3.3V。",
    "外部电源与USB可能同时插入时，必须设计电源选择/防回灌，不能依赖用户永远记得拔线。",
])
add_callout("范围控制", "车辆板是本项目第一优先级。遥控器板可在车辆板P1完成设计冻结后再进入PCB阶段，避免同时在两块板上重复犯首板错误。", "note")

page_break()
add_stage_intro("8", "ERC、封装绑定与设计同步", "在进入PCB前证明原理图逻辑和库映射完整。", "车辆/遥控器原理图、库文件", "Annotate结果、ERC报告、BOM、PCB同步文件", "无未解释ERC错误；每个实装位号都能找到封装")
add_heading("10.1 标注与属性检查", 2)
add_steps([
    "执行Annotate，确保车辆板使用U1/U2等，遥控器使用U101等；不要让两个工程混用位号序列。",
    "打开Property Editor，批量检查Reference、Value、PCB Footprint、DNP、Manufacturer、MPN。",
    "按BOM逐行核对，确保原理图没有遗漏机械孔、测试点和可选件。",
])
add_heading("10.2 ERC处理原则", 2)
add_bullets([
    "先运行Design Rules Check/ERC，再逐条分类：真实错误、符号引脚类型造成的合理警告、确实有意悬空。",
    "No ERC只放在经过解释的特定引脚/网络，禁止为了报告好看而全局屏蔽。",
    "重点检查：电源输入/输出冲突、未连接电源脚、重复位号、未连接连接器针脚、总线/跨页网络命名错误。",
    "保存ERC报告到manufacturing/vehicle/P1/reports，报告中每个剩余警告要写解释。",
])
add_heading("10.3 同步到PCB Editor", 2)
add_steps([
    "在OrCAD X Capture中选择设计根节点，执行PCB > Design Sync Setup，选择OrCAD/Allegro PCB Editor或OrCAD X Layout目标。",
    "使用PCB > New Layout/Update Layout创建或更新板文件。旧版本流程可能使用Tools > Create Netlist，再在PCB Editor中Import Logic。",
    "指定输出RC_CAR_VEHICLE_P1.brd和网表目录；确认PCB Footprint属性名被正确读取。",
    "阅读同步日志。任何Missing symbol、Missing padstack、Duplicate refdes必须解决后才能继续。",
    "在PCB Editor状态窗口核对元件数量、网络数量和未放置器件数量，与Capture和BOM一致。",
])
add_callout("设计同步纪律", "PCB阶段需要修改引脚或网络时，原则上返回Capture修改，再同步到PCB；不要同时在原理图和PCB各改一份，造成数据库分叉。", "risk")

page_break()
add_stage_intro("9", "板框、叠层和约束", "先用规则定义可制造边界，再开始布局。", "同步后的.brd、底盘孔位、板厂能力、电流数据", "板框/Keepin/安装孔、2层叠层、CSET规则", "规则有依据且已分配到网络；板框与底盘匹配")
add_heading("11.1 板框与机械对象", 2)
add_steps([
    "车辆板初始外形使用70×50 mm，最大不得超过80×60 mm。根据实测底盘孔位调整，不为缩小几毫米牺牲天线净空和操作空间。",
    "在Board Geometry / Outline绘制闭合板框；建立Route Keepin和Package Keepin，并与板边保留板厂要求间距。",
    "放置4个M3 NPTH 3.2 mm安装孔；安装孔周围按铜柱/垫片外径建立铜和器件避让。",
    "确认板框没有重复线、缝隙、自交或零长度段；板厂用于铣边的文件只保留一个清晰闭合轮廓。",
])
add_heading("11.2 叠层", 2)
add_table(["层", "用途", "原则"], [
    ("TOP", "元器件、关键信号、局部电源铜", "优先短路径；高电流用宽线或Shape"),
    ("BOTTOM", "连续GND参考和补充走线", "尽量保持地铜连续，减少被信号切割"),
    ("板材", "FR-4，1.6 mm，1 oz", "P1标准双层板；下单前与板厂确认"),
], [1.1, 2.1, 3.3])
add_heading("11.3 Constraint Manager", 2)
add_para("进入Tools > Constraint Manager。Cadence使用Constraint Set（CSET）集中定义物理和间距规则，再分配给网络。以下仅为P1起始建议，必须服从板厂能力和实测电流。")
add_table(["CSET", "对象", "起始建议", "最终确定依据"], [
    ("SIG", "GPIO/I2C/PWM", "线宽0.20-0.25 mm；间距≥0.20 mm", "板厂常规能力；优先0.25 mm便于首板"),
    ("LOGIC_PWR", "3V3/ESP_5V", "线宽≥0.50 mm", "负载电流和温升"),
    ("VBAT_PWR", "电池/VM/舵机", "优先2-3 mm宽铜或Shape，尽量不换层", "实测峰值电流、铜厚、长度、温升"),
    ("MOTOR_OUT", "AO/BO到电机", "优先1.5-2.5 mm或更宽，路径短", "启动/堵转电流和端子额定值"),
    ("EDGE", "铜到板边", "≥0.50 mm起始", "板厂铣边能力"),
], [1.05, 1.4, 2.15, 1.9], font_size=8.9)
add_callout("电流规则", "线宽表中的数字不是保证值。F1、开关、连接器、排针、铜箔和线束必须作为一条完整电流链共同核算。扩展到四电机前必须重新实测。", "risk")

add_stage_intro("10", "布局", "用器件位置决定电流路径、可装配性、射频和机械可靠性。", "板框、约束、器件、底盘方向", "完成布局、可操作性检查、初步3D/高度检查", "所有器件已放置；天线和高电流回路满足规则；模块可拔插")
add_heading("12.1 推荐布局顺序", 2)
add_steps([
    "先固定H1-H4安装孔、板框和所有板边连接器。电池、电机、舵机端子开口朝板外。",
    "将U1放在板边，天线端伸向外侧；天线下方和外侧不放铜、不走线、不放电池/电机/金属铜柱。",
    "将J1、F1、Q1、C1、U2/U3和J4-J7组成紧凑功率区，使电池到驱动到电机的去程和回程都短。",
    "将J3和C2放在舵机线入口附近，SERVO_VBAT和GND形成独立大电流分支。",
    "J2与D1靠近U1的5V入口，但与U1天线区保持距离。",
    "将低矮电阻、电容先放，再考虑高电解电容、模块拔插、USB插头和螺丝刀空间。",
    "最后放测试点，使探针在模块装上后仍能触达。",
])
add_heading("12.2 推荐区域关系", 2)
add_code("[板外/天线净空]  U1 SuperMini  [J2/逻辑电源]\n\n[J1/F1/Q1] -> [C1 + U2/U3功率区] -> [J4/J5/J6/J7板边]\n                       |\n                    [J3 + C2 舵机分支]\n\n四角：H1-H4；中间/板边空位：TP1-TP6")
add_heading("12.3 布局检查", 2)
add_checklist([
    "U1和U2/U3可垂直拔插，不被连接器、电容或铜柱阻挡。",
    "USB插头可插拔，线缆不会顶住高器件。",
    "螺丝端子可用螺丝刀操作，线材不会跨越天线。",
    "电解电容极性丝印装配后仍可见。",
    "电池、舵机和电机回流不必穿过U1区域。",
    "所有板边连接器开口方向正确，丝印不被器件遮挡。",
    "安装孔周围有足够铜柱和垫片净空。",
])

add_stage_intro("11", "布线与覆铜", "先完成电流回路，再布控制信号，最后形成连续接地参考。", "通过布局检查的.brd、CSET", "完成走线和Shape、初步DRC", "无未布网络；高电流路径满足规则；地回流连续")
add_heading("13.1 布线顺序", 2)
add_steps([
    "锁定或标记关键机械对象，先布VBAT_RAW/VBAT_PROTECTED/SERVO_VBAT/VM/MOTOR_OUT。",
    "高电流网络尽量在同一层用宽线或铜皮完成；避免颈缩、锐角、无意义换层和过孔串联。",
    "布REG_5V、ESP_5V和3V3_LOGIC，确保D1方向正确。",
    "布LEFT/RIGHT控制、STBY、SERVO_PWM、BAT_ADC和I2C等信号。PWM和ADC远离电机输出线。",
    "在TOP/BOTTOM铺GND Shape；重新计算/更新铜皮，检查孤岛、狭缝和回流被切断。",
    "在高电流回路附近使用足够铜连接，必要时增加地过孔，但天线Keepout内禁止过孔和铜。",
])
add_heading("13.2 关键回路", 2)
add_table(["回路", "正确做法", "错误做法"], [
    ("电池-驱动-电机", "正向和回流路径靠近、短、宽", "让电机电流绕板一周或经过U1地脚"),
    ("舵机", "J1保护后独立分支到J3/C2", "用ESP32 5V/3V3或细排针供电"),
    ("逻辑电源", "外置5V经隔离后到U1", "与USB 5V直接硬并联"),
    ("BAT_ADC", "分压/滤波靠近U1 ADC，远离PWM", "长线贴着电机输出走"),
    ("天线", "朝板外并保持完整禁布", "天线下铺地或放电池/铜柱"),
], [1.4, 2.55, 2.55])
add_heading("13.3 丝印和装配信息", 2)
add_bullets([
    "标板名、版本、日期、正负极、连接器功能、U2/U3用途、USB/天线方向。",
    "所有1脚、二极管A/K、电解正负、LED极性均可见。",
    "在J3同时标GND/PWR/SIG；在J2标IN+/GND/5V/GND；在电机端子标L/R和前/后。",
    "删除焊盘上的丝印；保持字符高度/线宽满足板厂规则。",
])

add_stage_intro("12", "DRC、评审与设计冻结", "证明PCB数据库、电气规则和机械规则均满足首板要求。", "完成走线的.brd、原理图、BOM、实物", "DRC报告、评审记录、P1冻结包", "DRC为0或每个豁免有书面理由；设计评审通过")
add_heading("14.1 PCB检查顺序", 2)
add_steps([
    "更新全部动态Shape，检查是否存在未连接、孤岛、细颈和意外热焊盘。",
    "运行在线/批量DRC，按Spacing、Physical、Same Net、Shape、Manufacturing分类处理。",
    "运行Database Check或等效数据库检查，解决零长度线、重复对象、坏图形等问题。",
    "查看Status，确认Unplaced Components=0、Unrouted Nets=0、DRC Errors=0（或有已批准豁免）。",
    "用3D Canvas/STEP检查器件高度、USB/插头/螺丝刀空间；没有STEP模型时用装配外框和实物尺寸人工复核。",
])
add_heading("14.2 Capture-PCB一致性复核", 2)
add_checklist([
    "Capture与PCB的位号、数量、封装和网络数量一致。",
    "U1/U2/U3 Pin-Pad映射再次逐脚复核。",
    "所有DNP器件在BOM和装配图中可识别。",
    "原理图修改均已重新同步，PCB没有孤立的手工逻辑修改。",
    "设计文件、库和固件GPIO基线一致。",
])
add_heading("14.3 设计评审重点", 2)
add_table(["评审域", "必须回答的问题"], [
    ("安全", "反接、过流、USB回灌和失联停车是否均有明确方案？"),
    ("电源", "最坏工况下电池、开关、保险丝、连接器、铜箔和线束是否有余量？"),
    ("封装", "所有自建封装是否1:1套孔？极性和方向是否复核？"),
    ("射频", "天线位置、铜禁布和机械遮挡是否合格？"),
    ("制造", "最小线宽/间距/孔径/环宽/板边间距是否满足板厂？"),
    ("装配", "器件能否按顺序焊接、插拔、测量和维修？"),
    ("测试", "无模块时能否测试电源？关键网络是否有测试点？"),
], [1.15, 5.35])
add_callout("冻结动作", "评审通过后复制为RC_CAR_VEHICLE_P1_RELEASE.brd并设为只读备份。后续生产输出只从该冻结文件生成。", "pass")

add_stage_intro("13", "生产文件输出", "从同一个冻结数据库生成完整、相互一致的制造和装配文件。", "P1_RELEASE.brd、冻结BOM、板厂能力", "Gerber/钻孔/BOM/坐标/图纸/报告/README", "输出层完整；文件名清晰；与冻结板一致")
add_heading("15.1 Gerber/Artwork", 2)
add_para("在传统PCB Editor中通常使用Manufacture > Artwork；在OrCAD X Layout中寻找Manufacturing/Export Gerber等效入口。建议使用Gerber RS-274X，并为每一层建立独立Film。")
add_table(["输出", "建议文件", "必须检查"], [
    ("Top Copper", "RC_CAR_P1_TOP.art", "全部顶层铜、焊盘和Shape"),
    ("Bottom Copper", "RC_CAR_P1_BOTTOM.art", "全部底层铜和地铜"),
    ("Top Solder Mask", "RC_CAR_P1_TSM.art", "焊盘开窗，无遗漏/过大"),
    ("Bottom Solder Mask", "RC_CAR_P1_BSM.art", "通孔和底层焊盘开窗"),
    ("Top Silkscreen", "RC_CAR_P1_TSS.art", "字符不压焊盘，极性清楚"),
    ("Bottom Silkscreen", "RC_CAR_P1_BSS.art", "仅在实际使用时输出"),
    ("Board Outline", "RC_CAR_P1_OUTLINE.art", "单一闭合外形、槽孔说明"),
    ("Top/Bottom Paste", "RC_CAR_P1_TPS/BPS.art", "只有需要钢网贴装时输出"),
], [1.35, 2.25, 2.9], font_size=8.8)
add_heading("15.2 钻孔", 2)
add_steps([
    "使用Manufacture > NC > NC Drill或等效命令输出Excellon钻孔。",
    "确认PTH和NPTH是否分文件；板厂如要求分开，分别输出并命名。",
    "输出钻孔表/Drill Legend，检查工具号、孔径和数量；M3成品孔应为3.2 mm NPTH。",
    "如果存在槽孔，按板厂要求输出NC Route或机械层说明，不能只在丝印画槽。",
])
add_heading("15.3 装配与管理文件", 2)
add_bullets([
    "BOM：Reference、Qty、Value、Manufacturer、MPN、Footprint、DNP、备注。",
    "坐标/CPL：位号、X、Y、Rotation、Side；模块载板以手焊为主仍建议输出供核对。",
    "装配图：器件外框、位号、极性、连接器方向和DNP标识。",
    "原理图PDF、PCB PDF、3D/STEP（可选）、DRC/ERC报告、生产README。",
    "ODB++或IPC-2581可作为板厂接受时的补充，但首板仍建议保留独立Gerber+钻孔复核流程。",
])
add_heading("15.4 发布目录", 2)
add_code("manufacturing/vehicle/P1_RELEASE/\n  gerber/\n  drill/\n  assembly/\n  bom/\n  reports/\n  RC_CAR_VEHICLE_P1_README.txt\n  RC_CAR_VEHICLE_P1_RELEASE.brd\n  RC_CAR_VEHICLE_P1_schematic.pdf")
add_callout("一次生成", "生产文件必须在同一次发布中从同一个冻结.brd生成。不要把昨天的钻孔、今天的Gerber和更早的BOM混在一个压缩包。", "risk")

add_stage_intro("14", "Gerber复核与PCB下单", "独立验证生产数据，而不是相信编辑器画面。", "发布目录、独立Gerber Viewer、板厂下单页面", "Gerber复核记录、最终压缩包、订单参数", "每层可视检查通过；钻孔与焊盘对中；订单参数一致")
add_heading("16.1 独立复核", 2)
add_steps([
    "用Gerbv、KiCad Gerber Viewer或板厂在线查看器打开所有Gerber和钻孔文件。不要只在PCB Editor里看。",
    "先单层查看铜、阻焊、丝印、外形和钻孔，再叠加查看对准关系。",
    "检查板框是否只有一条；钻孔是否位于焊盘中心；NPTH是否正确；阻焊是否覆盖/开窗正确。",
    "检查Top/Bottom是否镜像关系正确；丝印文字方向、极性、板名和版本是否可读。",
    "测量外形和安装孔中心距，与机械冻结表一致。",
    "检查文件中没有测试用图形、旧版本Logo、临时文字或意外铜皮。",
])
add_heading("16.2 推荐下单参数", 2)
add_table(["参数", "P1建议", "注意"], [
    ("层数", "2层", "TOP/BOTTOM"), ("板材/厚度", "FR-4 / 1.6 mm", "标准工艺"),
    ("铜厚", "1 oz", "高电流能力仍由线宽/温升决定"), ("表面处理", "HASL无铅或ENIG", "模块载板均可；按预算"),
    ("阻焊", "绿色", "首板检查方便"), ("丝印", "白色", "确保最小线宽/字高满足板厂"),
    ("外形", "约70×50 mm", "最终以Gerber和实测孔位为准"), ("数量", "5片或板厂最小量", "首板先小批验证"),
], [1.2, 2.0, 3.3])
add_checklist([
    "在线预览的层数和文件识别正确。",
    "板厂没有把NPTH识别为PTH或忽略槽孔。",
    "订单尺寸与查看器测量一致。",
    "提交压缩包中不存在多个版本的同名层。",
    "下单截图、订单号和上传文件归档。",
], "下单前最终确认")

add_stage_intro("15", "来料检查与焊接", "在通电前排除制造缺陷、错料和装配错误。", "PCB、BOM、实物器件、万用表、焊接工具", "来料记录、装配完成板、断电检查记录", "无可见短路/断路；极性和方向正确；电源网络不短路")
add_heading("17.1 裸板检查", 2)
add_checklist([
    "板名、版本、尺寸和安装孔位置正确。",
    "板边无严重毛刺、裂纹或分层。",
    "焊盘、过孔、阻焊和丝印完整，无明显偏位。",
    "用万用表确认VBAT与GND、5V与GND、3V3与GND没有硬短路。",
    "随机抽查关键网络连续性：J1-F1-Q1-U2 VM、J3电源、U1 5V、各电机端子。",
])
add_heading("17.2 推荐焊接顺序", 2)
add_steps([
    "先焊低矮SMD电阻、电容、二极管和保护器件；每焊完一个电源器件就检查极性和短路。",
    "焊测试点、排阻/小器件和低矮排针。",
    "焊螺丝端子、舵机座和模块排母；使用实际模块临时定位排母，避免歪斜，但不要带电焊接。",
    "最后焊高电解电容、开关和高器件，确认不会挡住模块和螺丝刀。",
    "清洁助焊剂残留，用放大镜检查连锡、虚焊、反向、漏焊和焊盘翘起。",
])
add_heading("17.3 模块未插入时的断电检查", 2)
add_table(["测量", "预期", "异常处理"], [
    ("VBAT_RAW-GND电阻", "非持续近0Ω", "查J1、F1、Q1、C1极性和焊桥"),
    ("REG_5V-GND电阻", "非持续近0Ω", "查J2/D1/电容"),
    ("3V3-GND电阻", "未插U1时应无硬短", "查接口和焊桥"),
    ("模块相邻针脚", "仅设计上应连的网络导通", "查排母连锡"),
    ("安装孔", "与铜网隔离（若设计为NPTH）", "查制造或封装设置"),
], [1.85, 2.1, 2.55])

add_stage_intro("16", "分阶段上电与固件", "用最小风险逐级验证电源、逻辑、驱动、舵机和无线。", "装配板、限流电源/电池、万用表、固件、车轮悬空支架", "各阶段电压/电流/温升记录、可运行整车", "每一级通过后再接下一级负载；异常立即断电")
add_callout("首次上电原则", "优先使用带限流功能的实验电源模拟电池并设置保守限流。若只能使用电池，串联适当保险保护并准备立即断电。车轮必须悬空，舵机拉杆首次可暂时断开。", "risk")
add_heading("18.1 阶段A：只验证外置5V模块", 2)
add_steps([
    "拔下U1、U2、U3，不接电机和舵机。",
    "给J1或模块输入供电，测VBAT_RAW/VBAT_PROTECTED。",
    "调节/确认外置模块输出为4.9-5.1 V，再连接J2输出到PCB。",
    "检查D1后ESP_5V的电压和极性；异常立即断电。",
])
add_heading("18.2 阶段B：安装U1", 2)
add_steps([
    "断电后插入SuperMini，方向与USB/天线丝印一致；不要同时插USB和外部5V。",
    "上电测U1 5V约5.0 V、3V3约3.3 V，观察启动电流和温升。",
    "通过USB烧录时关闭车辆电池，并断开外部模块OUT+到U1 5V的路径；烧录后拔USB再恢复。",
])
add_heading("18.3 阶段C：安装U2但不接电机", 2)
add_steps([
    "断电插入U2，确认VM/VCC方向；U3保持不装。",
    "上电测U2 VM≈电池电压、VCC≈3.3 V。",
    "运行测试固件，确认STBY在保护状态约0 V、解锁时约3.3 V；检查AIN/BIN/PWM逻辑。",
])
add_heading("18.4 阶段D：逐个连接电机", 2)
add_steps([
    "车轮悬空，先接一个后轮电机，以小PWM测试正反转和停止。",
    "确认无异常电流、复位、冒烟、过热和连接器发热后，再接第二个电机。",
    "保持通信层正油门=前进；当前软件使用MOTOR_A_DIRECTION=-1、MOTOR_B_DIRECTION=-1修正安装极性。",
])
add_heading("18.5 阶段E：连接舵机", 2)
add_steps([
    "先断开转向拉杆或让机构无机械顶死风险。",
    "确认J3为GND/SERVO_VBAT/SERVO_PWM，舵机红线不得接U1 5V/3V3。",
    "使用校准程序先到1380 us中位，再测试1120 us左端和1640 us右端；不得超出当前机械安全范围。",
    "连接拉杆后检查全行程，若顶死、抖动或电池压降立即缩小范围并排查机械阻力。",
])
add_heading("18.6 阶段F：无线与整车", 2)
add_bullets([
    "遥控器程序：software/arduino/tests/rc_car_sender_a_failsafe/rc_car_sender_a_failsafe.ino。",
    "车辆程序：software/arduino/tests/rc_car_vehicle_remote_motor_test/rc_car_vehicle_remote_motor_test.ino。",
    "通道6；A板MAC 44:B1:76:08:40:2C，B板MAC 44:B1:76:07:B4:C0。",
    "摇杆回中并松开按键后解锁；按键急停或失联后必须再次回中解锁。",
])
add_heading("18.7 上电记录", 2)
add_table(["阶段", "VBAT", "5V", "3V3", "静态/峰值电流", "结果/备注"], [
    ("A 外置5V", "____", "____", "—", "____", "____"),
    ("B U1", "____", "____", "____", "____", "____"),
    ("C U2无电机", "____", "____", "____", "____", "____"),
    ("D 双电机", "____", "____", "____", "____", "____"),
    ("E 舵机", "____", "____", "____", "____", "____"),
    ("F 整车", "____", "____", "____", "____", "____"),
], [0.9, 0.75, 0.7, 0.7, 1.45, 2.0], font_size=8.7)

add_stage_intro("17", "整车验收与P2闭环", "验证车辆在真实机械和负载条件下满足功能、安全和可靠性要求。", "完整车辆、遥控器、充满电池、测试场地", "验收记录、问题清单、P2变更说明", "所有必测项通过；高风险问题关闭；生产包归档")
add_heading("19.1 功能验收", 2)
add_checklist([
    "上电时车辆不突然移动，舵机回到安全中位。",
    "摇杆前/后与车辆前进/后退一致；左右转向一致。",
    "两个后轮在相同油门下方向一致，无明显单侧失速。",
    "按下急停后电机立即停止，松开并回中后才能重新解锁。",
    "关闭遥控器或屏蔽通信，100 ms失联保护有效。",
    "满油门、转向和无线同时工作时ESP32不复位。",
    "连续运行后TB6612、端子、保险丝、线束和PCB铜箔无异常发热。",
])
add_heading("19.2 机械与射频验收", 2)
add_checklist([
    "板体固定牢靠，线束不进入车轮或转向拉杆活动范围。",
    "USB、模块、端子和测试点仍可操作。",
    "天线未被电池、电机、金属铜柱或大电流线束遮挡。",
    "转向在1120-1640 us范围内不顶死，舵机无持续堵转声。",
    "不同方向和距离下ESP-NOW通信稳定，失联时安全停车。",
])
add_heading("19.3 首板问题分级", 2)
add_table(["等级", "定义", "例子", "处理"], [
    ("A 安全", "可能损坏器件/伤人/起火", "电源反接、过热、失联不停车", "立即停用；必须P2修正"),
    ("B 功能", "核心功能无法稳定实现", "电机掉电复位、针序错误", "飞线验证后P2修正"),
    ("C 制造", "可工作但难装配/难维修", "端子被挡、丝印不清", "记录并P2优化"),
    ("D 优化", "不影响P1使用", "板略大、测试点位置一般", "按价值排序"),
], [0.9, 1.8, 2.2, 1.6])
add_heading("19.4 P2变更流程", 2)
add_steps([
    "为每个问题记录现象、复现条件、测量数据、照片、原因和临时处理。",
    "先修改原理图和库，再同步PCB；不要只改Gerber或只在PCB上补丁。",
    "更新版本号、BOM、固件GPIO文档、测试步骤和生产README。",
    "重新执行受影响阶段的检查；电源/封装/板框变化时不得只做局部DRC。",
    "生成全新的P2_RELEASE目录，保留P1所有文件用于比较。",
])

add_heading("附录A：关键网络速查", 1)
add_table(["网络", "源", "负载", "布局/布线要求"], [
    ("VBAT_RAW", "J1", "F1/Q1", "极性清楚，入口短"),
    ("VBAT_PROTECTED", "保护后", "U2/U3 VM、J2、舵机分支", "大电流宽铜"),
    ("SERVO_VBAT", "保护后独立分支", "J3 Pin2/C2", "不得经过U1"),
    ("REG_5V", "J2 Pin3", "D1", "与USB隔离"),
    ("ESP_5V", "D1后", "U1 5V", "测试点可达"),
    ("3V3_LOGIC", "U1 3V3", "U2/U3 VCC、扩展", "仅逻辑"),
    ("MOTOR_STBY", "GPIO6", "U2/U3 STBY", "失联拉低"),
    ("SERVO_PWM", "GPIO21", "J3 Pin3", "远离电机线"),
    ("BAT_ADC", "分压中点", "GPIO0", "短、滤波、远离PWM"),
], [1.35, 1.6, 1.7, 1.85])

add_heading("附录B：发布包总检查表", 1)
add_checklist([
    "冻结.brd、.dsn、.olb、.pad、.dra、.psm全部备份。",
    "ERC报告无未解释严重问题；DRC/Database Check通过。",
    "Gerber层齐全且独立查看器复核通过。",
    "PTH/NPTH钻孔和槽孔文件正确。",
    "BOM、DNP、坐标和装配图版本一致。",
    "板名、版本、外形、孔位、连接器方向和极性正确。",
    "README列出层文件、叠层、板厚、铜厚和特殊说明。",
    "上传压缩包已归档并计算文件日期/版本。",
])

add_heading("附录C：常见问题定位", 1)
add_table(["现象", "优先检查", "不要先做"], [
    ("看得见封装但选不中", "Find过滤器与当前命令", "重装软件"),
    ("找不到封装", "psmpath、.psm名称、大小写/同名冲突", "随意复制到安装目录"),
    ("缺少焊盘", "padpath与封装调用的Padstack名称", "重新画整个封装"),
    ("鼠标行为异常", "按Esc或右键Done退出当前命令", "强制结束程序"),
    ("同步后缺元件", "PCB Footprint属性、同步日志、库路径", "在PCB手工补孤立器件"),
    ("电机启动复位", "电池压降、地回流、线径、C1/C2、连接器", "提高PWM前继续试"),
    ("舵机抖动", "供电、共地、机械阻力、信号线远离电机", "扩大脉宽范围"),
    ("急停后不运行", "摇杆是否完全回中并松开按键", "删除安全状态机"),
], [1.6, 3.15, 1.75], font_size=8.9)

add_heading("附录D：参考资料", 1)
add_bullets([
    "项目固定约束：H:\\RC-Car\\agent.md",
    "已验证接线：H:\\RC-Car\\docs\\RC-Car-V1_详细接线说明.md",
    "既有首板设计资料：H:\\RC-Car\\docs\\RC-Car-V1_Altium-Designer首版PCB打样步骤.md（电气与机械基线可迁移，具体操作以Cadence为准）",
    "Cadence OrCAD X Capture培训与流程：https://www.cadence.com/en_US/home/training/all-courses/85037.html",
    "Cadence OrCAD X Capture/PCB Editor教程：https://resources.pcb.cadence.com/documentation-page",
    "Cadence Constraint Manager说明：https://resources.academic.cadence.com/pcb-design-and-analysis-resources/2024-how-to-use-the-constraint-manager-in-orcad-x",
    "Cadence PCB封装流程：https://resources.academic.cadence.com/pcb-layout/2024-how-to-create-a-pcb-footprint-in-orcad-x-cadence",
    "Espressif ESP32-C3硬件设计指南：https://docs.espressif.com/projects/esp-hardware-design-guidelines/en/latest/esp32c3/pcb-layout-design.html",
    "Toshiba TB6612FNG官方资料：https://toshiba.semicon-storage.com/ap-en/semiconductor/product/motor-driver-ics/brushed-dc-motor-driver-ics/detail.TB6612FNG.html",
])
add_callout("最终原则", "先冻结、再建库；先验证封装、再画原理图；先规则、再布局布线；先独立检查生产文件、再下单；先分阶段上电、再接全部负载。", "pass")

# Final formatting controls
for p in doc.paragraphs:
    p.paragraph_format.widow_control = True
    if p.style.name.startswith("Heading"):
        p.paragraph_format.keep_with_next = True

# Document metadata
props = doc.core_properties
props.title = "RC-CAR-V1.0 Cadence完整设计与生产SOP"
props.subject = "从元器件建库、原理图、PCB到打样装车的完整流程"
props.author = "RC-CAR Project"
props.keywords = "Cadence, OrCAD X, PCB Editor, RC-Car, SOP"
props.comments = "基于RC-CAR-V1.0已验证架构与Cadence官方工作流整理。"

doc.save(DOCX)
print(DOCX)
